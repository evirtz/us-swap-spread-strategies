"""Read the supplied workbook observations without changing the source files."""
from datetime import date,datetime
import math
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl

RAW_COLUMNS = [
    "10yr swap spreads","30yr swap spreads","Move",
    "2yr_Basis_GBP_SOFR","5yr_Basis_GBP_SOFR",
    "2yr_Basis_EU_SOFR","5yr_Basis_EU_SOFR",
]
RATE_COLUMNS = ["SOFRRATE"] + [
    f"{kind}_{tenor}yrs" + ("_SOFR" if kind=="IRS" else "")
    for kind in ("UST","IRS") for tenor in (7,10,20,30)
]

def read_sofr(path,columns):
    book=openpyxl.load_workbook(Path(path),read_only=True,data_only=True)
    try:
        if "SOFR" not in book.sheetnames:
            raise ValueError(f"Missing SOFR sheet: {path}")
        rows=book["SOFR"].iter_rows(values_only=True)
        headers=next(rows)
        if headers[0]!="Date":
            raise ValueError(f"Expected Date as the first SOFR column: {path}")
        for name in columns:
            if headers.count(name)!=1:
                raise ValueError(f"Expected exactly one {name!r} column: {path}")
        positions=[headers.index(name) for name in columns]
        records=[]
        for row in rows:
            if not isinstance(row[0],(date,datetime)):
                continue
            values=[]
            for index in positions:
                value=row[index]
                values.append(float(value) if isinstance(value,(int,float))
                              and not isinstance(value,bool) and math.isfinite(value) else np.nan)
            records.append([pd.Timestamp(row[0])]+values)
    finally:
        book.close()
    frame=pd.DataFrame(records,columns=["Date"]+columns)
    if frame.empty or frame.Date.duplicated().any():
        raise ValueError(f"Missing or duplicate source dates: {path}")
    return frame.set_index("Date").sort_index()

def load_inputs(regression_path,rates_path):
    raw=read_sofr(regression_path,RAW_COLUMNS)
    rates=read_sofr(rates_path,RATE_COLUMNS).rename(columns={"SOFRRATE":"Floating"})
    if rates.Floating.dropna().empty:
        raise ValueError("SOFRRATE contains no saved numeric observations.")
    return raw,rates
