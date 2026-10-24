#!/usr/bin/env python3
"""Reproduce the latest SOFR gross monthly returns and verify them against Excel.

Dependencies: numpy, openpyxl.

Example (put the two Excel files beside this script, or pass their full paths):
    python reproduce_sofr_gross_returns.py --verify

This is the monthly CAPITAL-RETURNS stage of the analysis. It reads existing
signed market/carry/rolldown marks and entry Treasury notionals per DV01 from
Swap_Spread_Rolling_Results.xlsx. It does not refit the rolling regressions or
rebuild the original carry/rolldown estimates from the two raw workbooks.

Gross includes baseline Treasury SOFR funding already included in source carry.
It adds no bid/offer, repo haircut, swap margin costs, or idle-cash interest.
All workbook content is read as data; no workbook formulas or macros execute.
Input Excel files are never edited. Output CSV returns are decimals, not percent.
Blank CSV returns are distinguished by Status: Not started versus After capital
exhaustion. The monthly CSV includes cumulative values used by the Excel plots.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import numpy as np
import openpyxl

FAMILIES = tuple(
    f"SOFR_{maturity}Y_{currency}{tenor}Y"
    for maturity in (10, 30)
    for currency in ("GBP", "EUR")
    for tenor in (2, 5)
)
MODELS = tuple(f"{family}_R{r}" for family in FAMILIES for r in (1, 2, 3))
LEVERAGES = (1, 10, 20, 30)
STARTS = {1: date(2022, 2, 22), 2: date(2023, 2, 21), 3: date(2024, 2, 21)}
FINAL_CLOSE = date(2026, 9, 3)
INITIAL_CAPITAL = 100.0
REFERENCE_NAME = "SOFR_Gross_Monthly_Cumulative_Calendar_1x_10x_20x_30x.xlsx"


def require(condition: bool, message: str) -> None:
    """Raise an explicit error; validation remains active with python -O."""
    if not condition:
        raise ValueError(message)


def as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"Expected an Excel date, received {value!r}")


def finite_number(value, label: str) -> float:
    require(isinstance(value, (int, float)) and not isinstance(value, bool),
            f"Missing/non-numeric {label}: {value!r}")
    require(math.isfinite(value), f"Non-finite {label}: {value!r}")
    return float(value)


def month_start(d: date) -> date:
    return date(d.year, d.month, 1)


def month_range(first: date, last: date) -> tuple[date, ...]:
    dates = []
    cursor = month_start(first)
    while cursor <= month_start(last):
        dates.append(cursor)
        cursor = date(cursor.year + (cursor.month == 12), cursor.month % 12 + 1, 1)
    return tuple(dates)


MONTHS = month_range(min(STARTS.values()), FINAL_CLOSE)


def file_hash(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def table_rows(workbook, sheet_name: str, fields: tuple[str, ...]):
    """Source tables have their actual field names in Excel row 5."""
    require(sheet_name in workbook.sheetnames, f"Missing source sheet: {sheet_name}")
    iterator = workbook[sheet_name].iter_rows(min_row=5, values_only=True)
    headers = next(iterator)
    require(all(field in headers for field in fields),
            f"Missing required columns in {sheet_name}: {set(fields)-set(headers)}")
    positions = {field: headers.index(field) for field in fields}
    for row in iterator:
        if row[positions["Model"]] in MODELS:
            yield {field: row[i] for field, i in positions.items()}


def load_source(path: Path):
    """Read signed P&L changes and fixed trade attributes from the gross source."""
    daily = defaultdict(list)
    trades = defaultdict(dict)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for row in table_rows(workbook, "Daily SOFR", (
            "Model", "Date", "Marked trade", "Market change (bp)",
            "Carry change (bp)", "Roll change (bp)",
        )):
            daily[row["Model"]].append({
                "Date": as_date(row["Date"]), "Trade": row["Marked trade"] or None,
                "MarketBp": finite_number(row["Market change (bp)"], "market change"),
                "CarryBp": finite_number(row["Carry change (bp)"], "carry change"),
                "RolldownBp": finite_number(row["Roll change (bp)"], "roll change"),
            })
        for row in table_rows(workbook, "Trades", (
            "Model", "Trade ID", "Entry execution", "Exit / final mark", "Status",
            "Treasury notional per DV01", "Swap notional per DV01",
            "Market P&L (bp)", "Net carry (bp)", "Net rolldown (bp)",
        )):
            model, tid = row["Model"], row["Trade ID"]
            require(tid not in trades[model], f"Duplicate trade: {tid}")
            treasury = finite_number(row["Treasury notional per DV01"], "Treasury notional/DV01")
            swap = finite_number(row["Swap notional per DV01"], "swap notional/DV01")
            require(treasury > 0 and swap > 0, f"Nonpositive entry notional/DV01: {tid}")
            trades[model][tid] = {
                "EntryDate": as_date(row["Entry execution"]),
                "EndDate": as_date(row["Exit / final mark"]), "Status": row["Status"],
                "TreasuryNotionalPerDV01": treasury, "SwapNotionalPerDV01": swap,
                "MarketBp": finite_number(row["Market P&L (bp)"], "trade market P&L"),
                "CarryBp": finite_number(row["Net carry (bp)"], "trade carry"),
                "RolldownBp": finite_number(row["Net rolldown (bp)"], "trade rolldown"),
            }
    finally:
        workbook.close()

    require(set(daily) == set(MODELS) and set(trades) == set(MODELS),
            "The source does not contain all 24 selected rolling models.")
    trade_reconciliation_error = 0.0
    for model in MODELS:
        daily[model].sort(key=lambda row: row["Date"])
        rows = daily[model]
        dates = [row["Date"] for row in rows]
        require(len(dates) == len(set(dates)), f"Duplicate source dates: {model}")
        require(dates[0] == STARTS[int(model[-1])] and dates[-1] == FINAL_CLOSE,
                f"Unexpected return period for {model}: {dates[0]} to {dates[-1]}")
        by_trade = defaultdict(list)
        for row in rows:
            if row["Trade"]:
                require(row["Trade"] in trades[model], f"Unknown marked trade: {row['Trade']}")
                by_trade[row["Trade"]].append(row)
            else:
                require(all(abs(row[field]) < 1e-9 for field in ("MarketBp", "CarryBp", "RolldownBp")),
                        f"P&L without a marked trade: {model}, {row['Date']}")
        require(set(by_trade) == set(trades[model]), f"Missing trade marks: {model}")
        for tid, marks in by_trade.items():
            t = trades[model][tid]
            require(marks[0]["Date"] == t["EntryDate"] and marks[-1]["Date"] == t["EndDate"],
                    f"Trade date/mark mismatch: {tid}")
            for field in ("MarketBp", "CarryBp", "RolldownBp"):
                error = abs(math.fsum(row[field] for row in marks) - t[field])
                trade_reconciliation_error = max(trade_reconciliation_error, error)
                require(error < 1e-8, f"Daily marks do not reconcile to {field}: {tid}")
    return daily, trades, trade_reconciliation_error


def simulate_account(model: str, leverage: int, source: list[dict], lookup: dict):
    """Fix DV01 within a trade; resize from available capital at the next entry."""
    capital = INITIAL_CAPITAL
    sizes, entry_capitals = {}, {}
    records = []
    for row in source:
        opening = capital
        tid = row["Trade"]
        if tid:
            t = lookup[tid]
            if tid not in sizes:
                require(row["Date"] == t["EntryDate"] and opening > 0,
                        f"Invalid trade entry: {model}, {tid}")
                entry_capitals[tid] = opening
                # Treasury notional = leverage * entry capital. Treasury and
                # swap have equal money DV01; this size stays fixed in the trade.
                sizes[tid] = leverage * opening / t["TreasuryNotionalPerDV01"]
            money_dv01 = sizes[tid]
        else:
            money_dv01 = 0.0
        # Source changes already contain the long/short sign. Do not sign twice.
        parts = money_dv01 * np.array([row["MarketBp"], row["CarryBp"], row["RolldownBp"]])
        capital += float(parts.sum())
        records.append({
            "Model": model, "Leverage": leverage, "Date": row["Date"], "Trade": tid,
            "OpeningCapital": opening, "ClosingCapital": capital,
            "DailyReturn": capital / opening - 1,
            "MoneyDV01PerBp": money_dv01,
            "TreasuryEntryNotional": leverage * entry_capitals[tid] if tid else 0.0,
            "SwapEntryNotional": money_dv01 * lookup[tid]["SwapNotionalPerDV01"] if tid else 0.0,
            "MarketPnL": float(parts[0]), "CarryPnL": float(parts[1]), "RolldownPnL": float(parts[2]),
            "MarketBp": row["MarketBp"], "CarryBp": row["CarryBp"], "RolldownBp": row["RolldownBp"],
        })
        # Include the actual breach-day close, even if negative; never restart.
        if capital <= 0:
            break

    exhausted = capital <= 0
    require(all(row["ClosingCapital"] > 0 for row in records[:-1]), "Exhaustion stop failed")
    grouped = defaultdict(list)
    for row in records:
        if row["Trade"]:
            grouped[row["Trade"]].append(row)
    trade_capital, independently_compounded = [], INITIAL_CAPITAL
    for tid, marks in grouped.items():
        t = lookup[tid]
        completed = t["Status"] == "Completed" and marks[-1]["Date"] == t["EndDate"]
        status = "Completed" if completed else "Stopped: capital exhausted" if exhausted else "Open at final close"
        components = {field: math.fsum(row[field] for row in marks)
                      for field in ("MarketBp", "CarryBp", "RolldownBp")}
        trade_return = leverage * math.fsum(components.values()) / t["TreasuryNotionalPerDV01"]
        require(abs(independently_compounded - entry_capitals[tid]) < 1e-7,
                f"Trade opening capital mismatch: {tid}")
        independently_compounded *= 1 + trade_return
        require(abs(independently_compounded - marks[-1]["ClosingCapital"]) < 1e-7,
                f"Trade/daily capital mismatch: {tid}")
        trade_capital.append({
            "Model": model, "Leverage": leverage, "Trade": tid, "Status": status,
            "EntryDate": t["EntryDate"], "EndDate": marks[-1]["Date"],
            "EntryCapital": entry_capitals[tid], "ClosingCapital": marks[-1]["ClosingCapital"],
            "MoneyDV01PerBp": sizes[tid], "TradeReturn": trade_return, **components,
        })
    trade_capital_error = abs(independently_compounded - capital)
    require(trade_capital_error < 1e-7, f"Reinvestment does not reconcile: {model}, {leverage}")
    summary = {
        "Model": model, "Leverage": leverage, "FirstClose": source[0]["Date"],
        "LastClose": records[-1]["Date"], "PlannedLastClose": source[-1]["Date"],
        "Status": "Capital exhausted" if exhausted else "Full period",
        "ExhaustionDate": records[-1]["Date"] if exhausted else None,
        "FinalCapital": capital, "CumulativeReturn": capital / INITIAL_CAPITAL - 1,
        "TradeCapitalError": trade_capital_error,
    }
    return records, trade_capital, summary


def aggregate_account(records: list[dict], source: list[dict], summary: dict):
    """Monthly ratios, compounded cumulative returns, and eligible calendar years."""
    actual, planned = defaultdict(list), defaultdict(list)
    for row in records:
        actual[month_start(row["Date"])].append(row)
    for row in source:
        planned[month_start(row["Date"])].append(row)
    monthly, growth = [], 1.0
    for month in MONTHS:
        rows = actual[month]
        expected_rows = planned[month]
        status = ("Not started" if not expected_rows else "After capital exhaustion" if not rows
                  else "Capital exhausted" if rows[-1]["ClosingCapital"] <= 0 else "Available")
        item = {
            "Model": summary["Model"], "Leverage": summary["Leverage"], "Month": month,
            "Status": status, "MonthlyReturn": None, "CumulativeReturn": None,
            "OpeningCapital": None, "ClosingCapital": None,
            "MarketPnL": None, "CarryPnL": None, "RolldownPnL": None,
            "FirstClose": rows[0]["Date"] if rows else None, "LastClose": rows[-1]["Date"] if rows else None,
            "PartialMonth": bool(expected_rows and (month in (month_start(source[0]["Date"]), month_start(FINAL_CLOSE))
                                  or bool(rows) and len(rows) < len(expected_rows))),
            "IncludedSessions": len(rows), "PlannedSessions": len(expected_rows),
        }
        if rows:
            opening, closing = rows[0]["OpeningCapital"], rows[-1]["ClosingCapital"]
            parts = {field: math.fsum(row[field] for row in rows)
                     for field in ("MarketPnL", "CarryPnL", "RolldownPnL")}
            monthly_return = closing / opening - 1
            require(abs(opening + math.fsum(parts.values()) - closing) < 1e-7, "Monthly P&L does not reconcile")
            require(abs(math.prod(1 + row["DailyReturn"] for row in rows) - 1 - monthly_return) < 1e-9,
                    "Monthly return does not match daily compounding")
            growth *= 1 + monthly_return
            require(abs(growth - closing / INITIAL_CAPITAL) < 1e-9, "Cumulative compounding failed")
            item.update(MonthlyReturn=monthly_return, CumulativeReturn=growth - 1,
                        OpeningCapital=opening, ClosingCapital=closing, **parts)
        monthly.append(item)
    require(abs(growth - summary["FinalCapital"] / INITIAL_CAPITAL) < 1e-9, "Final capital does not reconcile")

    annual = []
    for year in range(MONTHS[0].year, FINAL_CLOSE.year + 1):
        rows = [row for row in monthly if row["Month"].year == year and row["MonthlyReturn"] is not None]
        status = "After capital exhaustion" if year > summary["LastClose"].year else "Not started"
        item = {"Model": summary["Model"], "Leverage": summary["Leverage"], "Year": year,
                "CalendarReturn": None, "FirstClose": None, "LastClose": None, "Status": status}
        if rows:
            status = "Full year"
            if year == summary["FirstClose"].year:
                status = f"From {summary['FirstClose'].isoformat()}"
            if year == summary["LastClose"].year:
                status = f"To {summary['LastClose'].isoformat()}"
                if summary["Status"] == "Capital exhausted":
                    status += "; capital exhausted"
            item.update(CalendarReturn=math.prod(1 + row["MonthlyReturn"] for row in rows) - 1,
                        FirstClose=rows[0]["FirstClose"], LastClose=rows[-1]["LastClose"], Status=status)
        annual.append(item)
    return monthly, annual


def verify_excel(reference: Path, monthly: list[dict], annual: list[dict]) -> dict:
    """Compare every output cell, including blanks/n.a., and Monthly Marks inputs.

    Reference layout: monthly blocks start at rows 7, 69, 131 and 193; columns
    B:D are R1:R3; cumulative values are Q:S; calendar values are G:I starting
    at rows 32, 94, 156 and 218. Model headers are checked before their values.
    These are the actual latest gross-file locations.
    """
    monthly_lookup = {(r["Model"], r["Leverage"], r["Month"]): r for r in monthly}
    annual_lookup = {(r["Model"], r["Leverage"], r["Year"]): r for r in annual}
    report = {"Status": "PASS", "Reference": str(reference.resolve()), "ReturnToleranceDecimal": 1e-9,
              "Counts": defaultdict(int), "MaxAbsoluteDifference": defaultdict(float), "Mismatches": []}

    def compare(actual, expected, category, cell, *, blank=False, unavailable=False, tolerance=1e-9):
        report["Counts"][category] += 1
        if blank:
            ok = actual is None or actual == ""
        elif unavailable:
            ok = actual == "n.a."
        elif isinstance(expected, date):
            ok = isinstance(actual, (date, datetime)) and as_date(actual) == expected
        elif isinstance(expected, (str, bool)):
            ok = actual == expected
        else:
            ok = isinstance(actual, (int, float)) and not isinstance(actual, bool) and math.isfinite(actual)
            if ok:
                difference = abs(float(actual) - float(expected))
                report["MaxAbsoluteDifference"][category] = max(report["MaxAbsoluteDifference"][category], difference)
                ok = difference <= tolerance
        if not ok:
            report["Mismatches"].append({"Cell": cell, "Category": category,
                "Excel": actual.isoformat() if isinstance(actual, (date, datetime)) else actual,
                "Python": expected.isoformat() if isinstance(expected, date) else expected,
                "ExpectedBlank": blank, "ExpectedNA": unavailable})

    workbook = openpyxl.load_workbook(reference, read_only=True, data_only=True)
    try:
        require(set(FAMILIES).issubset(workbook.sheetnames), "Missing reference model-family sheets")
        for family in FAMILIES:
            sheet_rows = list(workbook[family].iter_rows(values_only=True))
            for j, leverage in enumerate(LEVERAGES):
                first = 7 + j * 62
                band = sheet_rows[first - 3][0]
                require(band == f"{leverage}× monthly returns", f"Unexpected leverage section in {family}: {band!r}")
                for header_row, first_column, suffix in (
                    (first - 1, 1, ""),
                    (first - 1, 16, " cumulative"),
                    (31 + j * 62, 6, ""),
                ):
                    for R in (1, 2, 3):
                        column = first_column + R - 1
                        compare(sheet_rows[header_row - 1][column], f"R{R}{suffix}",
                                "ModelHeaders", f"{family}!{chr(65+column)}{header_row}")
                for i, month in enumerate(MONTHS):
                    excel_row = first + i
                    cells = sheet_rows[excel_row - 1]
                    require(as_date(cells[0]) == month, f"Month alignment error: {family}!A{excel_row}")
                    for R in (1, 2, 3):
                        model = f"{family}_R{R}"
                        expected = monthly_lookup[model, leverage, month]
                        value = expected["MonthlyReturn"]
                        prestart = expected["Status"] == "Not started"
                        compare(cells[R], value, "MonthlyReturns", f"{family}!{chr(65+R)}{excel_row}",
                                blank=prestart, unavailable=value is None and not prestart)
                        compare(cells[R+15], expected["CumulativeReturn"], "CumulativeReturns",
                                f"{family}!{chr(80+R)}{excel_row}", blank=value is None)
                for i, year in enumerate(range(MONTHS[0].year, FINAL_CLOSE.year + 1)):
                    excel_row = 32 + j * 62 + i
                    cells = sheet_rows[excel_row - 1]
                    require(cells[5] == year, f"Calendar-year alignment error: {family}!F{excel_row}")
                    for R in (1, 2, 3):
                        expected = annual_lookup[f"{family}_R{R}", leverage, year]
                        value = expected["CalendarReturn"]
                        prestart = expected["Status"] == "Not started"
                        compare(cells[R+5], value, "CalendarReturns", f"{family}!{chr(70+R)}{excel_row}",
                                blank=prestart, unavailable=value is None and not prestart)

        require("Monthly Marks" in workbook.sheetnames, "Missing Monthly Marks reference sheet")
        expected_keys = {key for key, r in monthly_lookup.items() if r["Status"] != "Not started"}
        seen_keys = set()
        for n, row in enumerate(workbook["Monthly Marks"].iter_rows(min_row=14, values_only=True), start=14):
            if row[0] is None:
                continue
            key = row[0], int(row[1]), as_date(row[2])
            require(key in expected_keys and key not in seen_keys, f"Unexpected/duplicate Monthly Marks key: {key}")
            seen_keys.add(key)
            expected = monthly_lookup[key]
            for column, field in ((3, "OpeningCapital"), (4, "MarketPnL"), (5, "CarryPnL"),
                                  (6, "RolldownPnL"), (7, "ClosingCapital"), (8, "MonthlyReturn"),
                                  (9, "FirstClose"), (10, "LastClose"), (11, "Status"),
                                  (12, "PartialMonth"), (13, "IncludedSessions"), (14, "PlannedSessions")):
                value = expected[field]
                compare(row[column], value, field, f"Monthly Marks!{chr(65+column)}{n}",
                        blank=value is None and column not in (7, 8),
                        unavailable=value is None and column in (7, 8),
                        tolerance=1e-7 if column in (3, 4, 5, 6, 7) else 1e-9)
        require(seen_keys == expected_keys, "Missing reference Monthly Marks records")
    finally:
        workbook.close()
    report["Counts"] = dict(report["Counts"])
    report["MaxAbsoluteDifference"] = dict(report["MaxAbsoluteDifference"])
    if report["Mismatches"]:
        report["Status"] = "FAIL"
    return report


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        for row in rows:
            writer.writerow({key: value.isoformat() if isinstance(value, date) else value
                             for key, value in row.items()})


def main() -> None:
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--source", type=Path, default=here / "Swap_Spread_Rolling_Results.xlsx")
    parser.add_argument("--reference", type=Path, default=here / REFERENCE_NAME)
    parser.add_argument("--output-dir", type=Path, default=here / "gross_returns_python")
    parser.add_argument("--verify", action="store_true", help="Compare every return and Monthly Marks record with the reference Excel")
    args = parser.parse_args()
    require(args.source.is_file(), f"Source workbook not found: {args.source}")
    if args.verify:
        require(args.reference.is_file(), f"Reference workbook not found: {args.reference}")
        require(args.source.resolve() != args.reference.resolve(), "Source and reference must be different files")
    before = {str(p.resolve()): file_hash(p) for p in (args.source, args.reference) if p.is_file()}
    source, trades, source_error = load_source(args.source)
    daily_capital, trade_capital, monthly, annual, accounts = [], [], [], [], []
    for model in MODELS:
        for leverage in LEVERAGES:
            records, ledger, summary = simulate_account(model, leverage, source[model], trades[model])
            months, years = aggregate_account(records, source[model], summary)
            daily_capital.extend(records)
            trade_capital.extend(ledger)
            monthly.extend(months)
            annual.extend(years)
            accounts.append(summary)
    require(len(accounts) == 96 and len(monthly) == 5376 and len(annual) == 480, "Unexpected output coverage")
    report = verify_excel(args.reference, monthly, annual) if args.verify else {"Status": "NOT REQUESTED"}
    report.update(Accounts=len(accounts), ModelFamilies=len(FAMILIES), RollingModels=len(MODELS),
                  MonthlyRows=len(monthly), EligibleMonthlyRows=sum(r["Status"] != "Not started" for r in monthly),
                  AvailableMonthlyReturns=sum(r["MonthlyReturn"] is not None for r in monthly),
                  SourceTradeReconciliationMaxBpError=source_error,
                  TradeCapitalReconciliationMaxError=max(a["TradeCapitalError"] for a in accounts),
                  CapitalExhaustionCounts={str(L): sum(a["Leverage"] == L and a["Status"] == "Capital exhausted" for a in accounts)
                                           for L in LEVERAGES},
                  InputSHA256=before)
    for filename, original_hash in before.items():
        require(file_hash(Path(filename)) == original_hash, f"Input file unexpectedly changed: {filename}")
    report["InputFilesUnchanged"] = True
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, rows in (("daily_capital", daily_capital), ("trade_capital", trade_capital),
                       ("monthly_returns", monthly), ("calendar_year_returns", annual), ("account_summary", accounts)):
        write_csv(args.output_dir / f"{name}.csv", rows)
    (args.output_dir / "verification_report.json").write_text(json.dumps(report, indent=2, allow_nan=False), encoding="utf-8")
    console_report = {key: value for key, value in report.items() if key not in ("Mismatches", "InputSHA256")}
    print(json.dumps(console_report, indent=2))
    print(f"Results saved in: {args.output_dir.resolve()}")
    if report["Status"] == "FAIL":
        raise SystemExit("Excel comparison FAILED; see verification_report.json for mismatched cells.")


if __name__ == "__main__":
    main()
