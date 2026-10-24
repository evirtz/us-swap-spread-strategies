"""Regression checks for model labels in the fixed Excel reference layout."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from scripts import reproduce_sofr_gross_returns as gross


# These are the headers in the saved gross reference workbook, at 1x leverage.
HEADER_SECTIONS = (
    ("Monthly", 6, 2, ("R1", "R2", "R3")),
    ("Cumulative", 6, 17, ("R1 cumulative", "R2 cumulative", "R3 cumulative")),
    ("Calendar", 31, 7, ("R1", "R2", "R3")),
)


class ExcelModelHeaderTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.reference = Path(temporary.name) / "reference.xlsx"
        self.family = gross.FAMILIES[0]
        self.enterContext(patch.object(gross, "FAMILIES", (self.family,)))
        self.workbook = openpyxl.Workbook()
        self.addCleanup(self.workbook.close)
        self.sheet = self.workbook.active
        self.sheet.title = self.family
        self.workbook.create_sheet("Monthly Marks")
        self.monthly = []
        self.annual = []

        # Inactive accounts need no numeric values or Monthly Marks records.
        # All leverage blocks and the full date grid still use the real layout.
        for block, leverage in enumerate(gross.LEVERAGES):
            offset = block * 62
            self.sheet.cell(5 + offset, 1, f"{leverage}\N{MULTIPLICATION SIGN} monthly returns")
            for _, header_row, first_column, labels in HEADER_SECTIONS:
                for column, label in enumerate(labels, start=first_column):
                    self.sheet.cell(header_row + offset, column, label)
            for row, month in enumerate(gross.MONTHS, start=7 + offset):
                self.sheet.cell(row, 1, month)
                for window in (1, 2, 3):
                    self.monthly.append({
                        "Model": f"{self.family}_R{window}", "Leverage": leverage,
                        "Month": month, "Status": "Not started",
                        "MonthlyReturn": None, "CumulativeReturn": None,
                    })
            years = range(gross.MONTHS[0].year, gross.FINAL_CLOSE.year + 1)
            for row, year in enumerate(years, start=32 + offset):
                self.sheet.cell(row, 6, year)
                for window in (1, 2, 3):
                    self.annual.append({
                        "Model": f"{self.family}_R{window}", "Leverage": leverage,
                        "Year": year, "Status": "Not started", "CalendarReturn": None,
                    })

    def verify(self):
        self.workbook.save(self.reference)
        return gross.verify_excel(self.reference, self.monthly, self.annual)

    def test_correct_headers_pass(self):
        report = self.verify()
        self.assertEqual(report["Status"], "PASS")
        self.assertEqual(report["Mismatches"], [])

    def test_swapped_labels_fail_in_every_section_and_leverage(self):
        for block, leverage in enumerate(gross.LEVERAGES):
            for section, row, column, _ in HEADER_SECTIONS:
                with self.subTest(section=section, leverage=leverage):
                    left = self.sheet.cell(row + block * 62, column)
                    right = self.sheet.cell(row + block * 62, column + 1)
                    original = left.value, right.value
                    try:
                        left.value, right.value = original[1], original[0]
                        report = self.verify()
                        self.assertEqual(report["Status"], "FAIL")
                        self.assertEqual(
                            {(item["Category"], item["Cell"]) for item in report["Mismatches"]},
                            {("ModelHeaders", f"{self.family}!{cell.coordinate}")
                             for cell in (left, right)},
                        )
                    finally:
                        left.value, right.value = original

    def test_missing_or_unknown_labels_fail_in_each_section(self):
        for section, row, column, _ in HEADER_SECTIONS:
            for replacement in (None, "R4"):
                with self.subTest(section=section, label=replacement):
                    cell = self.sheet.cell(row + 3 * 62, column + 2)
                    original = cell.value
                    try:
                        cell.value = replacement
                        report = self.verify()
                        self.assertEqual(report["Status"], "FAIL")
                        self.assertEqual(
                            [(item["Category"], item["Cell"]) for item in report["Mismatches"]],
                            [("ModelHeaders", f"{self.family}!{cell.coordinate}")],
                        )
                    finally:
                        cell.value = original

    def test_incorrect_return_still_fails_with_valid_headers(self):
        self.sheet["B7"] = 0.25
        report = self.verify()
        self.assertEqual(report["Status"], "FAIL")
        self.assertEqual(
            [(item["Category"], item["Cell"]) for item in report["Mismatches"]],
            [("MonthlyReturns", f"{self.family}!B7")],
        )


if __name__ == "__main__":
    unittest.main()
